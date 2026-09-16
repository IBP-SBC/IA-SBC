"""
nucleo/extraccion.py
====================
Convierte cualquier archivo soportado en TEXTO PLANO buscable.

PRINCIPIO: limpiar sí, interpretar no. Acá no se resume, no se
reescribe y no se "mejora" el contenido: se extrae lo que el archivo
dice, marcando de qué página/hoja/diapositiva salió para poder citarlo.

Si un formato no se puede leer (falta la librería, el PDF es una imagen
escaneada), se devuelve un AVISO CLARO en vez de un texto vacío. Un
documento que entra vacío es peor que uno rechazado: aparece en la lista
y el usuario cree que la IA lo conoce.
"""
from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path


class ResultadoExtraccion:
    """Texto extraído + diagnóstico honesto de cómo salió."""

    def __init__(self, texto: str = "", paginas: int = 0,
                 aviso: str = "", ok: bool = True):
        self.texto = texto
        self.paginas = paginas      # páginas / hojas / diapositivas
        self.aviso = aviso
        self.ok = ok and bool(texto.strip())


def _limpiar(texto: str) -> str:
    """Espacios y saltos de línea razonables, sin tocar el contenido."""
    t = texto.replace("\r\n", "\n").replace("\r", "\n")
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


# ════════════════════════════════════════════════════════════════════
# EXTRACTORES POR FORMATO
# ════════════════════════════════════════════════════════════════════

def _pdf(datos: bytes) -> ResultadoExtraccion:
    try:
        from pypdf import PdfReader
    except Exception:
        return ResultadoExtraccion(aviso="Falta la librería 'pypdf' en el entorno.", ok=False)
    try:
        lector = PdfReader(io.BytesIO(datos))
        partes, vacias = [], 0
        for i, pag in enumerate(lector.pages, start=1):
            try:
                txt = pag.extract_text() or ""
            except Exception:
                txt = ""
            if txt.strip():
                partes.append(f"[Página {i}]\n{txt}")
            else:
                vacias += 1
        total = len(lector.pages)
        aviso = ""
        if vacias and total:
            aviso = (f"{vacias} de {total} páginas no tienen texto extraíble "
                     "(probablemente son imágenes escaneadas); esas no se indexan.")
        return ResultadoExtraccion(_limpiar("\n\n".join(partes)), total, aviso)
    except Exception as e:
        return ResultadoExtraccion(aviso=f"No se pudo leer el PDF: {e}", ok=False)


def _docx(datos: bytes) -> ResultadoExtraccion:
    try:
        import docx  # python-docx
    except Exception:
        return ResultadoExtraccion(aviso="Falta la librería 'python-docx'.", ok=False)
    try:
        doc = docx.Document(io.BytesIO(datos))
        partes = [p.text for p in doc.paragraphs if p.text.strip()]
        # Las tablas de un Word suelen traer lo importante (matrices,
        # cronogramas). Se aplanan fila por fila separadas por ' | '.
        for tabla in doc.tables:
            for fila in tabla.rows:
                celdas = [c.text.strip() for c in fila.cells]
                if any(celdas):
                    partes.append(" | ".join(celdas))
        return ResultadoExtraccion(_limpiar("\n".join(partes)), 1)
    except Exception as e:
        return ResultadoExtraccion(aviso=f"No se pudo leer el Word: {e}", ok=False)


def _pptx(datos: bytes) -> ResultadoExtraccion:
    try:
        from pptx import Presentation
    except Exception:
        return ResultadoExtraccion(aviso="Falta la librería 'python-pptx'.", ok=False)
    try:
        pres = Presentation(io.BytesIO(datos))
        partes = []
        for i, slide in enumerate(pres.slides, start=1):
            textos = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                    textos.append(shape.text_frame.text.strip())
            if textos:
                partes.append(f"[Diapositiva {i}]\n" + "\n".join(textos))
        return ResultadoExtraccion(_limpiar("\n\n".join(partes)), len(pres.slides))
    except Exception as e:
        return ResultadoExtraccion(aviso=f"No se pudo leer la presentación: {e}", ok=False)


def _xlsx(datos: bytes) -> ResultadoExtraccion:
    try:
        import pandas as pd
    except Exception:
        return ResultadoExtraccion(aviso="Falta pandas en el entorno.", ok=False)
    try:
        hojas = pd.read_excel(io.BytesIO(datos), sheet_name=None, dtype=str)
    except Exception as e:
        return ResultadoExtraccion(aviso=f"No se pudo leer el Excel: {e}", ok=False)
    partes, recortadas = [], 0
    for nombre, df in hojas.items():
        if df is None or df.empty:
            continue
        df = df.fillna("")
        # Techo por hoja: una tabla de 20.000 filas no aporta más que sus
        # primeras 500 para responder preguntas, y sí llena el índice.
        if len(df) > 500:
            df = df.head(500)
            recortadas += 1
        filas = [" | ".join(str(c) for c in df.columns)]
        for _, fila in df.iterrows():
            valores = [str(v).strip() for v in fila.tolist()]
            if any(valores):
                filas.append(" | ".join(valores))
        partes.append(f"[Hoja: {nombre}]\n" + "\n".join(filas))
    aviso = (f"{recortadas} hoja(s) se recortaron a las primeras 500 filas."
             if recortadas else "")
    return ResultadoExtraccion(_limpiar("\n\n".join(partes)), len(hojas), aviso)


def _csv(datos: bytes) -> ResultadoExtraccion:
    # Los exports de la SBC vienen en utf-8 o latin-1 y con ';' o ','.
    texto = None
    for cod in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            texto = datos.decode(cod)
            break
        except Exception:
            continue
    if texto is None:
        return ResultadoExtraccion(aviso="No se pudo decodificar el CSV.", ok=False)
    try:
        muestra = texto[:4000]
        sep = csv.Sniffer().sniff(muestra, delimiters=",;\t|").delimiter
    except Exception:
        sep = ";" if muestra.count(";") > muestra.count(",") else ","
    filas = [l for l in texto.splitlines() if l.strip()][:2000]
    filas = [" | ".join(c.strip() for c in l.split(sep)) for l in filas]
    aviso = "Se indexaron las primeras 2.000 filas." if len(filas) >= 2000 else ""
    return ResultadoExtraccion(_limpiar("\n".join(filas)), 1, aviso)


def _texto_plano(datos: bytes) -> ResultadoExtraccion:
    for cod in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return ResultadoExtraccion(_limpiar(datos.decode(cod)), 1)
        except Exception:
            continue
    return ResultadoExtraccion(aviso="No se pudo decodificar el archivo.", ok=False)


def _html(datos: bytes) -> ResultadoExtraccion:
    base = _texto_plano(datos)
    if not base.ok:
        return base
    t = re.sub(r"(?is)<(script|style).*?</\1>", " ", base.texto)
    t = re.sub(r"(?s)<[^>]+>", "\n", t)
    t = (t.replace("&nbsp;", " ").replace("&amp;", "&")
          .replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"'))
    return ResultadoExtraccion(_limpiar(t), 1)


def _json(datos: bytes) -> ResultadoExtraccion:
    base = _texto_plano(datos)
    if not base.ok:
        return base
    try:
        obj = json.loads(base.texto)
        return ResultadoExtraccion(
            _limpiar(json.dumps(obj, ensure_ascii=False, indent=2)), 1)
    except Exception:
        return base   # si no es JSON válido, se indexa como texto


_MAPA = {
    ".pdf": _pdf, ".docx": _docx, ".pptx": _pptx,
    ".xlsx": _xlsx, ".xlsm": _xlsx, ".csv": _csv,
    ".txt": _texto_plano, ".md": _texto_plano,
    ".html": _html, ".htm": _html, ".json": _json,
}


# ════════════════════════════════════════════════════════════════════
# COHERENCIA ENTRE EXTENSIÓN Y CONTENIDO
# ════════════════════════════════════════════════════════════════════
# Caso REAL encontrado en los documentos de la SBC: varios archivos
# llamados .pdf son por dentro archivos de Office (empiezan con 'PK',
# la firma de un ZIP) y un .docx era texto plano. Vienen así de
# conversiones y descargas de Drive.
#
# La app NO los "arregla" adivinando: los rechaza diciendo exactamente
# qué son y cómo renombrarlos. Si la app leyera un .pdf como Word por su
# cuenta, el día que el contenido no cuadre nadie sabría por qué.

def _tipo_dentro_del_zip(datos: bytes) -> str:
    """Mira qué hay dentro de un ZIP para poder nombrar el formato real."""
    try:
        import zipfile
        with zipfile.ZipFile(io.BytesIO(datos)) as z:
            nombres = z.namelist()
        if any(n.startswith("word/") for n in nombres):
            return ".docx"
        if any(n.startswith("xl/") for n in nombres):
            return ".xlsx"
        if any(n.startswith("ppt/") for n in nombres):
            return ".pptx"
        # Caso frecuente: un PDF que alguien "convirtió" y quedó como un
        # paquete de imágenes, una por página. Ahí no hay texto que leer.
        # Umbral 40% (no 50%) porque estos paquetes suelen traer además un
        # manifest y archivos sueltos: en un caso real fueron 48 imágenes
        # de 97 entradas, y con 50% se habría escapado.
        imagenes = [n for n in nombres
                    if n.lower().endswith((".jpg", ".jpeg", ".png", ".tif"))]
        if len(imagenes) >= 3 and len(imagenes) >= len(nombres) * 0.4:
            return ".imagenes"
    except Exception:
        pass
    return ""


def detectar_tipo_real(datos: bytes) -> str:
    """Formato real según los primeros bytes. '' si no se reconoce."""
    if datos[:5] == b"%PDF-":
        return ".pdf"
    if datos[:2] == b"PK":
        return _tipo_dentro_del_zip(datos) or ".zip"
    return ""


def revisar_coherencia(nombre: str, datos: bytes) -> str:
    """
    Devuelve un mensaje si la extensión NO corresponde al contenido, o
    cadena vacía si todo está bien.
    """
    ext = Path(nombre).suffix.lower()
    real = detectar_tipo_real(datos)
    binarios = {".pdf", ".docx", ".xlsx", ".xlsm", ".pptx"}

    if ext in binarios and not real:
        return (f"'{nombre}' dice ser {ext} pero por dentro parece texto "
                "plano, no un archivo de ese tipo. Abrilo, verificá qué es "
                "y renombralo con la extensión correcta (.txt, .md o .html).")
    if real and ext in binarios and real != ext:
        if real == ".imagenes":
            return (f"'{nombre}' no es un {ext}: por dentro son las páginas "
                    "convertidas a imágenes. De ahí no sale texto. Subí el "
                    "PDF original, o pasale OCR primero.")
        if real == ".zip":
            return (f"'{nombre}' dice ser {ext} pero por dentro es un archivo "
                    "comprimido (ZIP). Descomprimilo y subí los documentos "
                    "que tenga adentro.")
        return (f"'{nombre}' dice ser {ext} pero por dentro es un {real}. "
                f"Renombralo a {real} y volvé a subirlo. No lo leo como "
                f"{real} por mi cuenta: si después algo no cuadra, nadie "
                "sabría de dónde salió.")
    return ""


def extraer(nombre: str, datos: bytes) -> ResultadoExtraccion:
    """
    Punto de entrada único. Elige el extractor por la EXTENSIÓN del
    nombre, después de comprobar que el contenido corresponde.
    """
    ext = Path(nombre).suffix.lower()
    fn = _MAPA.get(ext)
    if fn is None:
        return ResultadoExtraccion(
            aviso=f"No sé leer archivos '{ext}'. Formatos soportados: "
                  + ", ".join(sorted(_MAPA)), ok=False)

    problema = revisar_coherencia(nombre, datos)
    if problema:
        return ResultadoExtraccion(aviso=problema, ok=False)

    res = fn(datos)
    if not res.ok and not res.aviso:
        res.aviso = ("El archivo no dejó texto extraíble. Si es un PDF "
                     "escaneado, hay que pasarle OCR antes de subirlo.")
    return res
